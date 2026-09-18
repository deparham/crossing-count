"""Page 1 for someone who reads nothing else, in the PowerPoint and the PDF alike."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from pptx import Presentation

from crossing_count import sampling, validation
from crossing_count.report_pdf import build_pdf
from crossing_count.report_pptx import build_report, card_value, identity_rows

PEAK = {"role": "peak", "rank": 1, "length_min": 15, "day": "2026-08-22", "direction": "both",
        "level": "busy", "seed": 7}


def texts(path: Path) -> list[str]:
    return ["\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            for slide in Presentation(str(path)).slides]


def test_the_headline_says_which_way_and_by_how_many_people() -> None:
    wh = validation.window_headline
    assert wh("RetailNext", "people coming in", 41, 48) == (
        "RetailNext counted 7 fewer people coming in than were verified (41 against 48): "
        "an undercount of 14.6%.")
    assert wh("RetailNext", "people coming in", 52, 48).endswith("an overcount of 8.3%.")
    assert wh("RetailNext", "people coming in", 48, 48) == (
        "RetailNext counted the same number of people coming in as were verified (48).")
    few = wh("RetailNext", "people going out", 9, 11)  # no percentage on fewer than 30
    assert "2 fewer" in few and few.endswith("too few crossings for a percentage.")
    assert "%" not in few
    # unclear crossings: none of them real, or all of them
    assert "between 1 fewer and 1 more" in wh("RetailNext", "people", 12, 11, 2)
    assert "up to 2 fewer" in wh("RetailNext", "people", 11, 11, 2)
    assert "between 2 and 4 more" in wh("RetailNext", "people", 15, 11, 2)
    assert "-2.0% to +0.0%" in wh("RetailNext", "people", 48, 48, 1)
    opening = sampling.when(PEAK)
    assert wh("RetailNext", "people", 41, 48, 0, opening).startswith(
        "In the busiest 15 minutes of 22/08/2026, RetailNext counted 7 fewer")


def test_sampling_in_a_few_words() -> None:
    assert sampling.label(None).startswith("chosen by hand")
    assert sampling.label(PEAK).startswith("peak: the busiest 15 minutes")
    assert "seed 7" in sampling.label({**PEAK, "role": "control", "level": "normal"})
    assert sampling.when({**PEAK, "rank": 2}).startswith("In the 2nd busiest")
    assert sampling.when(None, "11:15 to 11:30 on 22/08/2026") == (
        "From 11:15 to 11:30 on 22/08/2026, ")
    assert sampling.when(None) == ""


def _data(tmp_path: Path) -> dict[str, Any]:
    img = np.full((720, 1280, 3), 90, np.uint8)
    cv2.imwrite(str(tmp_path / "frame.jpg"), img)
    thumbs = []
    for n in range(4):
        cv2.imwrite(str(tmp_path / f"t{n}.jpg"), img[:400, :400])
        thumbs.append({"path": str(tmp_path / f"t{n}.jpg"), "label": f"{n + 1} · IN"})
    return {
        "store_name": "Lismore", "store_code": "RW-128", "location": "Entrance",
        "report_date": "25/08/2026", "captured_date": "22/08/2026", "time_range": "11:15-11:30",
        "headline": [validation.window_headline("RetailNext", "people coming in", 41, 48, 0,
                                                sampling.when(PEAK))],
        "identity": {"validation_id": None, "sampling": sampling.label(PEAK),
                     "specification": "Ground Truth Specification v1.4", "gold_set": "none yet",
                     "footage": "clean: no RetailNext marks"},
        "caveats": ["RetailNext's numbers are Lismore's total: its entrances together."],
        "directions": [{"key": "in", "label": "Traffic In", "verified": 48, "unsure": 0,
                        "system": 41, "rate": True, "accuracy": 85.4}],
        "complete": True, "incomplete": [],
        "scope": sampling.scope(PEAK), "sample_notes": ["Traffic In: from 48 verified."],
        "frame": str(tmp_path / "frame.jpg"), "frame_caption": "RW-128-PB1 · 11:22:10",
        "crossings": [{"n": n, "time": "11:16:00", "camera": "RW-128-PB1",
                       "direction": "Traffic In", "found": "Detected, confirmed"}
                      for n in range(1, 49)],
        "thumbs": thumbs, "method": ["Footage: x.mp4", "Checked by a person"],
    }


def test_page_one_leads_with_the_result_and_what_it_rests_on(tmp_path: Path) -> None:
    data = _data(tmp_path)
    assert identity_rows(data)[0] == ("Validation ID",
                                      "draft: no ID until the validation is finalised")
    assert identity_rows({}) == []
    assert card_value(data["directions"][0], True) == ("SENSOR ACCURACY", "85.4%")
    assert card_value({"verified": 11, "system": 9, "rate": False}, True) == ("DIFFERENCE", "−2")
    pages = texts(build_report(data, tmp_path / "r.pptx"))
    cover = pages[0]
    assert cover.index("RetailNext counted 7 fewer") < cover.index("VERIFIED")  # words first
    assert "Validation ID: draft" in cover and "Sampling: peak: the busiest" in cover
    assert "Ground truth: Ground Truth Specification v1.4" in cover
    assert "Footage: clean" in cover and "Lismore's total" in cover
    assert "RW-128\n" not in cover  # the code line gives way to the result
    assert "RW-128-PB1 · 11:22:10" in cover  # room for the picture on page 1
    data["sample_notes"] = [f"Note {k}: " + "so many words " * 12 for k in range(6)]
    pages = texts(build_report(data, tmp_path / "full.pptx"))
    # no room left on page 1 for a useful picture: it has a page of its own, before the snapshots
    assert "RW-128-PB1 · 11:22:10" not in pages[0]
    assert "Validation frame" in pages[-2] and "RW-128-PB1 · 11:22:10" in pages[-2]
    data.update(identity={**data["identity"], "validation_id": "CC-VAL-2026-AB12-000001"})
    assert "Validation ID: CC-VAL-2026-AB12-000001" in texts(build_report(data,
                                                                          tmp_path / "f.pptx"))[0]


def test_the_pdf_says_the_same(tmp_path: Path) -> None:
    data = _data(tmp_path)
    raw = build_pdf(data, tmp_path / "r.pdf", compress=False).read_bytes()
    assert raw.startswith(b"%PDF")
    pages = len(re.findall(rb"/Type /Page\b(?!s)", raw))
    assert pages >= 4  # the result, the method and crossings, the snapshots
    for words in (b"RetailNext counted 7 fewer", b"Validation ID", b"Ground Truth Specification",
                  b"Lismore's total", b"How this was counted", b"Crossing snapshots",
                  b"Page 1 of %d" % pages):
        assert words in raw.replace(b"\\'", b"'"), words
    data.update(complete=False, incomplete=["2 of 5 stretches were not watched."], headline=[])
    raw = build_pdf(data, tmp_path / "i.pdf", compress=False).read_bytes()
    assert b"Validation incomplete: 2 of 5 stretches" in raw

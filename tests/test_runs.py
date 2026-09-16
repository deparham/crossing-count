"""Finalised validations (an ID, a locked folder with checksums, new versions) and the audit log."""

from __future__ import annotations

import json
import re
import stat
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from crossing_count import auditlog, runs
from crossing_count.webapp import Setup
from crossing_count.wizard import Wizard, WizardError
from crossing_count.wizard_app import create_wizard_app


def test_the_audit_log_shows_any_edit(tmp_path: Path) -> None:
    for k in range(3):
        auditlog.append("answer", user="Alex", after=k, root=tmp_path)
    v = auditlog.verify(tmp_path)
    assert v["intact"] and v["entries"] == 3 and v["head"] == auditlog.head(tmp_path)
    p = auditlog.path(tmp_path)
    lines = p.read_text().splitlines()
    p.write_text("\n".join([lines[0], lines[1].replace('"after": 1', '"after": 7'), lines[2]]) + "\n")
    changed = auditlog.verify(tmp_path)
    assert not changed["intact"] and changed["broken_at"] == 2
    p.write_text("\n".join([lines[0], lines[2]]) + "\n")  # one removed
    assert auditlog.verify(tmp_path)["broken_at"] == 2


@pytest.fixture
def reported(two_tile_video: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Wizard:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path))
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("manual")
    w.set_marks(True)
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": False}])
    w.set_direction("in")
    dur = float(w.state["duration_s"])
    w.manual_add("CN-9-PB1", round(dur * 0.3, 2), "in")
    w.manual_watched("CN-9-PB1", 0.0, dur)
    w.manual_done()
    w.set_sensor({"in": 2})
    w.set_store(name="Test store", code="CN-9", operator="Alex")
    w.make_report(None)
    return w


def test_a_finalised_validation_is_kept_and_locked(reported: Wizard, tmp_path: Path) -> None:
    run = reported.finalise()
    vid = run["validation_id"]
    assert re.fullmatch(r"CC-VAL-\d{4}-[0-9A-F]{4}-000001", vid)
    folder = Path(run["folder"])
    names = {p.name for p in folder.iterdir()}
    assert {"manifest.json", "result.json", "intervals.csv", "crossings.csv",
            "decisions.json"} <= names and any(n.endswith(".pptx") for n in names)
    assert any(n.endswith(".pdf") for n in names)
    # the finalised report says its ID on page 1; the draft it was made from said draft
    (final,) = [folder / n for n in names if n.endswith(".pptx")]
    cover = "\n".join(sh.text_frame.text for sh in Presentation(str(final)).slides[0].shapes
                      if sh.has_text_frame)
    assert f"Validation ID: {vid}" in cover
    draft = Presentation(reported.state["report"]["path"]).slides[0].shapes
    assert any("Validation ID: draft" in sh.text_frame.text for sh in draft if sh.has_text_frame)
    m = json.loads((folder / "manifest.json").read_text())
    assert m["validation_id"] == vid and m["footage"]["sha256"] and m["software"]["crossingcount"]
    assert m["system_under_test"]["numbers"]["sensor"] == {"in": 2}
    assert m["ground_truth"]["specification"] == "1.2" and m["hardware"]["cpus"]
    assert m["audit_log_head"] != auditlog.GENESIS and set(m["files"]) == names - {"manifest.json"}
    assert "counted" in (folder / "crossings.csv").read_text()
    assert not (folder / "result.json").stat().st_mode & stat.S_IWUSR  # read-only
    assert runs.verify(folder, tmp_path) == {"intact": True, "changed": [], "missing": [],
                                             "audit_log_ok": True}
    for change in (lambda: reported.manual_add("CN-9-PB1", 1.0, "in"),
                   lambda: reported.set_rules("exclude", "count"),
                   lambda: reported.make_report(None)):
        with pytest.raises(WizardError, match="finalised as"):
            change()
    actions = {e["action"] for e in auditlog.recent(200, tmp_path)}
    assert {"hand_count_added", "counting_finished", "report_made", "finalised"} <= actions
    assert auditlog.verify(tmp_path)["intact"]
    (folder / "result.json").chmod(0o644)
    (folder / "result.json").write_text("{}")  # tampering is found
    assert runs.verify(folder, tmp_path)["changed"] == ["result.json"]


def test_a_correction_is_a_new_version(reported: Wizard, tmp_path: Path) -> None:
    first = reported.finalise()
    reported.new_version()
    assert reported.state["supersedes"] == first["validation_id"] and reported.state["report"] is None
    with pytest.raises(WizardError, match="Make the report first"):
        reported.finalise()
    reported.set_store(report_date="15/09/2026")
    reported.make_report(None)
    second = reported.finalise()
    assert second["validation_id"].endswith("000002")
    m = json.loads((Path(second["folder"]) / "manifest.json").read_text())
    assert m["supersedes"] == first["validation_id"]
    assert runs.verify(Path(first["folder"]), tmp_path)["intact"]  # the first stays as it was
    assert [r["validation_id"] for r in runs.listing(tmp_path)] == [second["validation_id"],
                                                                   first["validation_id"]]


def test_the_runs_page_answers(reported: Wizard, tmp_path: Path) -> None:
    client = TestClient(create_wizard_app(tmp_path / "sites", tmp_path, folders=[tmp_path]))
    assert client.get("/runs/").status_code == 200
    got = client.get("/api/runs").json()
    assert got["runs"] == [] and got["audit"]["intact"] and got["audit"]["entries"] > 0
    assert client.post("/api/runs/reveal", json={"id": "../../etc"}).status_code == 404

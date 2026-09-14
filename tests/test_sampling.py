"""Sampling: which windows are validated is a stated, recorded rule, and every result says
what it describes (docs/GROUND_TRUTH_SPECIFICATION.md, section 9)."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from crossing_count import retailnext as rn
from crossing_count import sampling
from crossing_count import validation as v
from crossing_count.webapp import Setup
from crossing_count.wizard import Wizard


def _hhmm(m: int) -> str:
    return f"{m // 60:02d}:{m % 60:02d}"


def _day(outs: list[int]) -> list[dict[str, Any]]:
    return [{"start": _hhmm(600 + 15 * k), "finish": _hhmm(615 + 15 * k), "in": 1, "out": o,
             "validity": "complete"} for k, o in enumerate(outs)]


# 10:00-18:00 with two cameras (0.5 camera-hours per 15 minutes): a quiet morning, a busy
# lunchtime (140 an hour), a normal afternoon (60 an hour), a quiet evening
ROWS = _day([5] * 8 + [70] * 4 + [30] * 12 + [8] * 8)


def plan(mode: str = "peak", seed: int = 7, rows: list[dict[str, Any]] = ROWS,
         length: int = 15) -> dict[str, Any]:
    return sampling.plan(rn.windows(rows, length, "out"), length_min=length, direction="out",
                         cameras=2, mode=mode, seed=seed, day="2026-09-12")


def test_peak_windows_come_with_a_control_window_of_lower_traffic() -> None:
    p = plan()
    peaks = [w for w in p["chosen"] if w["role"] == "peak"]
    assert [w["start"] for w in peaks] == ["12:00", "12:15", "12:30"]  # the busiest, as before
    assert [w["start"] for w in rn.busiest(ROWS, 15, "out")] == ["12:00", "12:15", "12:30"]
    assert {w["level"] for w in peaks} == {"busy"}
    (control,) = [w for w in p["chosen"] if w["role"] == "control"]
    assert control["level"] == "normal" and "13:00" <= control["start"] < "16:00"
    assert (p["mode"], p["seed"], len(p["candidates"])) == ("peak", 7, 32)
    assert p == plan()  # the same seed draws the same control window
    quiet = plan(rows=_day([3] * 12))
    assert [w["role"] for w in quiet["chosen"]] == ["peak"] * 3
    assert "No control window" in quiet["notes"][0]


def test_stratified_and_random_windows_are_drawn_repeatably() -> None:
    s = plan("stratified", seed=3)
    assert sorted(w["level"] for w in s["chosen"]) == ["busy", "normal", "quiet"]
    assert "heavy traffic" in s["notes"][0]  # a level the day does not have, said
    r = plan("random", seed=11, length=30)
    assert r["chosen"] == plan("random", seed=11, length=30)["chosen"] and len(r["chosen"]) == 3
    starts = sorted(int(w["start"][:2]) * 60 + int(w["start"][3:]) for w in r["chosen"])
    assert all(b - a >= 30 for a, b in zip(starts, starts[1:], strict=False))  # never overlapping
    with pytest.raises(sampling.SamplingError, match="peak, stratified, random"):
        plan("quietest")


def test_every_result_says_what_it_describes() -> None:
    p = plan()
    peak = sampling.record(p, "12:00", "12:15")
    assert (peak["role"], peak["rank"], peak["level"], len(peak["candidates"])) == ("peak", 1, "busy", 32)
    text = sampling.scope(peak)
    assert "peak trading" in text and "not of the whole day" in text
    assert "12/09/2026" in text and "busy traffic" in text
    control = next(w for w in p["chosen"] if w["role"] == "control")
    said = sampling.scope(sampling.record(p, control["start"], control["until"]))
    assert "control window" in said and "seed 7" in said and "not peak trading" in said
    hand = sampling.record(p, "10:00", "10:15")
    assert hand["role"] == "chosen by hand" and "this period only" in sampling.scope(hand)
    assert "this period (11:30–11:45 on 12/09/2026) only" in sampling.scope(
        None, "11:30–11:45 on 12/09/2026")
    brief = sampling.brief(peak)
    assert brief is not None and "candidates" not in brief and brief["window"]["start"] == "12:00"


def test_the_error_is_broken_down_by_traffic_level() -> None:
    heavy = {"interval": "a", "direction": "out", "truth": 60, "system": 54, "covered_s": 900,
             "cameras": 1}  # 240 an hour
    normal = {"interval": "b", "direction": "out", "truth": 20, "system": 20, "covered_s": 900,
              "cameras": 1}  # 80 an hour
    lv = v.by_level([heavy, normal])
    assert list(lv) == ["Normal traffic", "Heavy traffic"]
    assert (lv["Heavy traffic"]["bias_pct"], lv["Normal traffic"]["bias_pct"]) == (-10.0, 0.0)
    assert v.level_sentence(lv, "RetailNext") == (
        "At normal traffic (1 interval, 20 verified crossings) RetailNext counted exactly as "
        "many; at heavy traffic (1 interval, 60 verified crossings) RetailNext counted 10.0% "
        "too few.")
    vals = [{"id": "p1", "store": "S1", "camera_hours": 0.25, "rows": [heavy],
             "sampling": {"role": "peak"}},
            {"id": "c1", "store": "S1", "camera_hours": 0.25, "rows": [normal],
             "sampling": {"role": "control"}},
            {"id": "p2", "store": "S2", "camera_hours": 0.25, "rows": [heavy],
             "sampling": {"role": "peak"}}]
    s = v.summary(vals)
    assert s["scope"] == {"peak": 2, "control": 1} and s["stores_without_control"] == ["S2"]
    assert set(s["by_store_traffic"]["S1"]) == {"Normal traffic", "Heavy traffic"}
    text = v.headline(s, "RetailNext")
    assert "Sampled: 2 peak windows and 1 control window." in text
    assert "No control window yet at S2" in text


def _texts(path: Path) -> list[str]:
    return ["\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            for slide in Presentation(str(path)).slides]


def test_the_window_is_kept_from_download_to_report_and_manifest(
        two_tile_video: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path))
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("manual")
    w.set_marks(True)
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": False}])
    w.state["marks"] = False
    w.set_direction("out")
    rec = sampling.record(plan(), "12:00", "12:15")
    w.use_download({"code": "CN-9", "name": "Test store", "cameras": ["CN-9-PB1"], "sampling": rec})
    assert w.state["sampling"] == rec and "sampling" not in w.state["retailnext"]
    dur = float(w.state["duration_s"])
    w.manual_add("CN-9-PB1", round(dur * 0.3, 2), "out")
    w.manual_watched("CN-9-PB1", 0.0, dur)
    w.manual_done()
    w.set_sensor({"out": 2})
    w.set_store(name="Test store", operator="Alex")
    cover = _texts(w.make_report(None))[0]
    assert "Sampling: peak trading" in cover and "not of the whole day" in cover
    r = w.result()
    assert r["sampling"]["role"] == "peak" and "candidates" not in r["sampling"]
    assert r["traffic"]["level"] in v.TRAFFIC_NAMES and r["by_level"]
    run = w.finalise()
    m = json.loads((Path(run["folder"]) / "manifest.json").read_text())
    assert (m["sampling"]["mode"], m["sampling"]["seed"], m["sampling"]["window"]["level"]) == (
        "peak", 7, "busy")
    assert len(m["sampling"]["candidates"]) == 32 and len(m["sampling"]["chosen"]) == 4
    assert m["traffic"] == r["traffic"]

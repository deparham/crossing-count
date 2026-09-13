"""The count wizard: footage list, progress, checking crossings, and the PowerPoint report."""

from __future__ import annotations

import json
import sys
import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from crossing_count.report_pptx import A4_H, A4_W, build_report
from crossing_count.webapp import Setup
from crossing_count.wizard import (
    POSSIBLE_REASONS,
    Progress,
    Wizard,
    WizardError,
    list_videos,
    mark_twins,
    prompt_priority,
)
from crossing_count.wizard_app import PPTX, create_wizard_app


def copy_outputs(src: Path) -> Callable[[Wizard], list[list[str]]]:
    """Stands in for gate.py + detect.py: copies finished outputs into the wizard's run folder."""
    def commands(w: Wizard) -> list[list[str]]:
        code = (f"import shutil; shutil.copytree({str(src)!r}, {str(w.run_dir)!r}, "
                f"dirs_exist_ok=True, ignore=shutil.ignore_patterns('review', 'manual', "
                f"'wizard', 'export')); print('  gating 00:00:30.00 / 00:01:00.00  (5.0x realtime)')")
        return [[sys.executable, "-c", code]]
    return commands


def wait(w: Wizard) -> None:
    for _ in range(400):
        if w.job_status()["status"] != "running":
            return
        time.sleep(0.05)
    raise AssertionError("the count did not finish")


def texts(path: Path) -> list[str]:
    pages = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts.extend(c.text for row in shape.table.rows for c in row.cells)
        pages.append("\n".join(parts))
    return pages


def test_footage_list_shows_only_videos(tmp_path: Path) -> None:
    for name in ("a.mp4", "B.MOV", "notes.txt", ".hidden.mp4"):
        (tmp_path / name).write_bytes(b"x")
    assert {v["name"] for v in list_videos([tmp_path, tmp_path / "missing"])} == {"a.mp4", "B.MOV"}


def test_progress_is_read_from_gate_and_detect_output() -> None:
    p = Progress(cameras=2)
    p.feed("  gating 00:07:30.00 / 00:15:00.00  ( 75.0x realtime)")
    assert p.stage == "gate" and p.overall() == pytest.approx(2.5)
    p.feed("  CN-1: derotated detector, 0 static objects ignored")
    p.feed("    tracking  50.0% of active time ( 1.5x realtime)")
    assert p.camera == "CN-1" and p.overall() == pytest.approx(5 + 47.5 * 0.5)
    p.feed("  CN-2: derotated detector, 1 static objects ignored")
    assert p.camera == "CN-2" and p.overall() == pytest.approx(5 + 47.5)
    p.feed("objc[42]: Class AVFFrameReceiver is implemented in both ...")
    assert not any(line.startswith("objc") for line in p.log)


def test_possible_misses_most_often_real_come_first() -> None:
    assert {"no_filter", "returned_same_track", "outward_no_mask"} <= set(POSSIBLE_REASONS)
    assert "uturn_no_mask" not in POSSIBLE_REASONS  # measured: never real
    assert sorted(["lost", "pending_expired", "no_filter"], key=prompt_priority) == [
        "no_filter", "pending_expired", "lost"]


def test_one_person_on_two_cameras_is_flagged() -> None:
    items: list[dict[str, Any]] = [
        {"id": "a", "kind": "counted", "picture": 0, "camera": "A", "direction": "in", "t": 10.0},
        {"id": "b", "kind": "counted", "picture": 1, "camera": "B", "direction": "in", "t": 11.2},
        {"id": "c", "kind": "counted", "picture": 1, "camera": "B", "direction": "out", "t": 11.0},
        {"id": "d", "kind": "possible", "picture": 0, "camera": "A", "direction": "in", "t": 9.0},
    ]
    mark_twins(items)
    assert "twin" not in items[0] and items[1]["twin"]["id"] == "a"
    assert "twin" not in items[2] and items[3]["twin"]["id"] == "a"


def test_count_check_and_report(two_tile_video: dict[str, Any], review_run_dir: Path,
                                tmp_path: Path) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path,
               copy_outputs(review_run_dir))
    s = Setup(w.video, two_tile_video["dir"])
    w.set_cameras(s.existing(), [t.as_dict() for t in s.tiles])
    assert [c["sensor"] for c in w.state["cameras"]] == ["CAM-A", "CAM-B"]
    assert w.state["store"]["code"] == "SYN"
    with pytest.raises(WizardError):
        w.start()  # no direction chosen yet
    w.set_direction("in")
    w.set_sensor({"in": 3})
    w.start()
    wait(w)
    assert w.job_status()["status"] == "done"
    items = w.check_items()
    counted = [i for i in items if i["kind"] == "counted"]
    assert counted and all(i["direction"] == "in" for i in items)
    for i in items:
        w.answer(i["id"], "yes" if i is counted[0] else "no")
    w.set_watch("skipped")
    c = w.counts()
    assert c["verified"]["in"] == 1 and c["confirmed"] == 1 and c["checked"]
    assert c["accuracy"]["in"]["error"] == 2  # RetailNext 3 against 1 verified
    with pytest.raises(WizardError, match="store name"):
        w.make_report()
    w.set_store(name="Lismore", code="SYN-1", report_date="25/08/2026")
    pages = texts(w.make_report())
    assert "Lismore SYN-1 Traffic System" in pages[0] and "VERIFIED COUNT" in pages[0]
    assert "SYSTEM COUNT" in pages[0] and "ACCURACY" in pages[0] and "25/08/2026" in pages[0]
    assert f"Page 1 of {len(pages)}" in pages[0] and "Crossing details" in pages[1]
    assert "Detected, confirmed" in pages[1]
    stats = w.state["job"]["stats"]
    assert stats["app_version"] and stats["video_s"] == pytest.approx(60.0, abs=1)
    assert "Crossings proposed by the" in pages[1] and "Report made with CrossingCount" in pages[1]
    again = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)  # resumes
    assert again.counts()["verified"]["in"] == 1 and again.state["report"]


def test_unwatched_movement_and_unsure_answers(two_tile_video: dict[str, Any],
                                              review_run_dir: Path, tmp_path: Path,
                                              monkeypatch: pytest.MonkeyPatch) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path,
               copy_outputs(review_run_dir))
    s = Setup(w.video, two_tile_video["dir"])
    w.set_cameras(s.existing(), [t.as_dict() for t in s.tiles])
    w.set_direction("in")
    w.set_sensor({"in": 1})
    w.start()
    wait(w)
    stretch = {"id": "u9", "camera": "CAM-A", "picture": 0, "start": 10.0, "end": 25.0}
    monkeypatch.setattr(w, "watch_ranges", lambda: [stretch])
    items = w.check_items()
    counted = [i for i in items if i["kind"] == "counted"]
    for i in items:
        w.answer(i["id"], "no")
    w.answer(counted[0]["id"], "unsure")
    with pytest.raises(WizardError):
        w.answer(counted[0]["id"], "maybe")
    w.set_watch("skipped")
    c = w.counts()
    assert c["checked"] and c["status"] == "incomplete" and c["unwatched_s"] == 15.0
    assert c["incomplete"] == [("1 of 1 stretches of movement near the line that the tool could "
                                "not explain were not watched (15 s of footage).")]
    assert c["verified"]["in"] == 0 and c["unsure"]["in"] == 1  # unsure is never counted
    assert c["accuracy_range"]["in"] == [100.0, 100.0]  # right if the unsure one was real
    w.set_store(name="Lismore", code="SYN-1")
    pages = texts(w.make_report())
    assert "INCOMPLETE" in pages[0] and "Validation incomplete: 1 of 1" in pages[0]
    assert "0–1" in pages[0]  # the verified count, both ways the unsure crossing could go
    assert "VALIDATION INCOMPLETE" in pages[1] and "Unclear to the checker" in pages[1]
    w.mark_watched("u9")
    assert w.counts()["status"] == "complete" and w.counts()["incomplete"] == []


def test_a_group_is_counted_as_its_people(two_tile_video: dict[str, Any], review_run_dir: Path,
                                          tmp_path: Path) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path,
               copy_outputs(review_run_dir))
    s = Setup(w.video, two_tile_video["dir"])
    w.set_cameras(s.existing(), [t.as_dict() for t in s.tiles])
    w.set_direction("in")
    w.set_sensor({"in": 3})
    w.start()
    wait(w)
    items = w.check_items()
    first = next(i for i in items if i["kind"] == "counted")
    for i in items:
        w.answer(i["id"], "no")
    with pytest.raises(WizardError, match="1 to 9"):
        w.answer(first["id"], "yes", people=0)
    w.answer(first["id"], "yes", people=3)
    c = w.counts()
    assert c["verified"]["in"] == 3 and c["confirmed"] == 1
    assert (c["groups"], c["group_people"]) == (1, 3)
    assert c["accuracy"]["in"]["accuracy_pct"] == 100.0  # RetailNext 3 against 3 people
    w.answer(first["id"], None)
    assert w.state["people"] == {} and w.counts()["verified"]["in"] == 0
    w.answer(first["id"], "yes", people=3)
    w.set_watch("skipped")
    w.set_store(name="Lismore", code="SYN-1")
    pages = texts(w.make_report())
    assert "1–3" in pages[1] and "group of 3" in pages[1]
    assert "groups crossing together: 3 people" in pages[1]


def test_a_check_is_logged_and_kept_when_the_count_runs_again(
        two_tile_video: dict[str, Any], review_run_dir: Path, tmp_path: Path) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path,
               copy_outputs(review_run_dir))
    s = Setup(w.video, two_tile_video["dir"])
    w.set_cameras(s.existing(), [t.as_dict() for t in s.tiles])
    w.set_direction("in")
    w.set_sensor({"in": 2})
    w.set_store(name="Lismore", code="SYN-1", operator="Pat")
    w.start()
    wait(w)
    first, *rest = w.check_items()
    for i in rest:
        w.answer(i["id"], "no")
    w.answer(first["id"], "yes", people=2)
    w.answer(first["id"], None)
    w.answer(first["id"], "yes")
    w.add("CAM-A", 12.0, "in")
    w.remove_added(0)
    w.set_watch("skipped")
    log = w.state["decisions"]
    assert log[0]["action"] == "run" and [d["action"] for d in log[-6:]] == [
        "answer", "undo", "answer", "add", "remove_added", "watch"]
    assert log[-6]["people"] == 2 and log[-6]["item"]["t"] == first["t"]
    assert all(d["by"] == "Pat" and d["at"] for d in log)
    report = w.make_report()
    assert "Pat checked every crossing" in texts(report)[1]

    with pytest.raises(WizardError, match="Running it again starts a new check"):
        w.start()
    w.start(confirm=True)
    wait(w)
    assert w.state["answers"] == {} and [d["action"] for d in w.state["decisions"]] == ["run"]
    (h,) = w.history()
    assert h["answers"] == len(rest) + 1 and h["verified"] == {"in": 1} and h["by"] == "Pat"
    kept = json.loads(Path(h["file"]).read_text())
    assert kept["answers"][first["id"]] == "yes" and Path(kept["report_copy"]).is_file()
    assert w.state["decisions"][0]["kept"] == h["file"]


def test_retailnext_per_interval_and_per_camera(two_tile_video: dict[str, Any],
                                                review_run_dir: Path, tmp_path: Path) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path,
               copy_outputs(review_run_dir))
    s = Setup(w.video, two_tile_video["dir"])
    w.set_cameras(s.existing(), [t.as_dict() for t in s.tiles])
    w.state["clock_start"] = "2026-09-12T11:59:40+10:00"  # the 60 s clip spans two intervals
    w.set_direction("in")
    assert [(i["key"], i["full"]) for i in w.intervals()] == [("11:45", False), ("12:00", False)]
    with pytest.raises(WizardError, match="Enter RetailNext's number for 12:00"):
        w.set_sensor(intervals={"11:45": {"in": 1}})
    with pytest.raises(WizardError, match="not one of this footage's 15-minute intervals"):
        w.set_sensor(intervals={"09:00": {"in": 1}})
    w.set_sensor(intervals={"11:45": {"in": 1}, "12:00": {"in": 3}}, cameras={"CAM-A": {"in": 2}})
    assert w.state["sensor"] == {"in": 4}  # the total is the intervals' sum
    w.start()
    wait(w)
    for i in w.check_items():
        w.answer(i["id"], "yes")
    comp = w.comparison()
    assert sum(i["verified"]["in"] for i in comp["intervals"]) == w.counts()["verified"]["in"]
    cams = {c["camera"]: c for c in comp["cameras"]}
    assert cams["CAM-A"]["accuracy"]["in"]["sensor"] == 2 and cams["CAM-B"]["accuracy"] is None
    w.set_watch("skipped")
    w.set_store(name="Lismore", code="SYN-1")
    report = w.make_report()
    page = texts(report)[1]
    assert "Traffic by 15 minutes" in page and "Only partly in the footage" in page
    assert "CAM-A" in page and "In verified" in page
    assert any(sh.has_chart for sh in Presentation(str(report)).slides[1].shapes)
    w.set_sensor({"in": 5})  # a plain total replaces the intervals
    assert w.state["sensor_intervals"] == {} and w.state["sensor"]["in"] == 5


def test_report_layout(tmp_path: Path) -> None:
    img = np.full((480, 640, 3), 90, np.uint8)
    cv2.imwrite(str(tmp_path / "frame.jpg"), img)
    cv2.imwrite(str(tmp_path / "logo.png"), np.full((124, 302, 3), 255, np.uint8))
    thumbs = []
    for n in range(13):
        p = tmp_path / f"t{n}.jpg"
        cv2.imwrite(str(p), img[:400, :400])
        thumbs.append({"path": str(p), "label": f"{n + 1} · IN"})
    data: dict[str, Any] = {
        "store_name": "Lismore", "store_code": "RW-128", "location": "Entrance",
        "report_date": "25/08/2026", "captured_date": "22/08/2026", "time_range": "11:15-11:30",
        "directions": [{"key": "in", "label": "Traffic In", "verified": 11, "system": 10,
                        "accuracy": 90.9},
                       {"key": "out", "label": "Traffic Out", "verified": 8, "system": 8,
                        "accuracy": 100.0}],
        "frame": str(tmp_path / "frame.jpg"), "frame_caption": "RW-128-PB1 · 11:22:10",
        "crossings": [{"n": n, "time": "11:16:00", "camera": "RW-128-PB1",
                       "direction": "Traffic In", "found": "Detected, confirmed"}
                      for n in range(1, 60)],
        "thumbs": thumbs, "method": ["Footage: x.mp4", "Checked by a person"],
    }
    out = build_report(data, tmp_path / "r.pptx", tmp_path / "logo.png")
    prs = Presentation(str(out))
    assert (prs.slide_width, prs.slide_height) == (A4_W, A4_H)  # A4 portrait
    pages = texts(out)
    assert len(pages) == 6  # cover, 3 of details (59 rows), 2 of snapshots (13)
    assert "ACCURACY IN" in pages[0] and "90.9%" in pages[0] and "100%" in pages[0]
    assert "Crossing details (continued)" in pages[2] and "Crossing snapshots" in pages[4]
    assert "Page 6 of 6" in pages[5]

    data["directions"][0].update(unsure=2, accuracy_range=[81.8, 90.9])
    data.update(complete=False, incomplete=["2 of 5 stretches were not watched (1 min 0 s)."])
    cover = texts(build_report(data, tmp_path / "r2.pptx", None))[0]
    assert "INCOMPLETE" in cover and "90.9%" not in cover and "11–13" in cover
    assert "Validation incomplete: 2 of 5 stretches" in cover
    data["complete"] = True  # everything watched, two crossings still unclear
    cover = texts(build_report(data, tmp_path / "r3.pptx", None))[0]
    assert "81.8–90.9%" in cover and "INCOMPLETE" not in cover

    data["breakdown"] = {
        "dirs": [{"key": "in", "label": "Traffic In"}],
        "intervals": [{"label": f"11:{m:02d} - 11:{m + 15:02d}", "partial": m == 0,
                       "verified": {"in": 4}, "sensor": {"in": 5},
                       "accuracy": {"in": {"error": 1, "error_pct": 25.0}}} for m in (0, 15, 30)],
        "cameras": [{"label": c, "partial": False, "verified": {"in": 6}, "sensor": None,
                     "accuracy": None} for c in ("PB1", "R2")],
        "notes": ["Largest difference: 11:00 - 11:15."]}
    out = build_report(data, tmp_path / "r4.pptx", None)
    page = texts(out)[1]
    assert "Traffic by 15 minutes" in page and "+1 (+25%)" in page and "(part)" in page
    assert "PB1" in page and "Largest difference" in page and len(texts(out)) == 7
    assert any(sh.has_chart for sh in Presentation(str(out)).slides[1].shapes)


def test_wizard_page_api(two_tile_video: dict[str, Any], review_run_dir: Path,
                         tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path / "home"))  # never the real settings
    client = TestClient(create_wizard_app(two_tile_video["dir"], tmp_path,
                                          [two_tile_video["dir"]], None,
                                          copy_outputs(review_run_dir)))
    assert "Count wizard" in client.get("/").text
    assert client.get("/api/state").status_code == 409
    names = [v["name"] for v in client.get("/api/videos").json()["videos"]]
    assert names == ["two_cameras.mp4"]
    st = client.post("/api/open", json={"path": str(two_tile_video["video"])}).json()
    assert len(client.get(st["draw_url"] + "api/info").json()["pictures"]) == 2
    assert client.get("/api/drawn").json()["pictures"][0]["configs"][0]["sensor"] == "CAM-A"
    client.post("/api/cameras")
    client.post("/api/direction", json={"direction": "in"})
    client.post("/api/sensor", json={"values": {"in": 2}})
    assert client.post("/api/run").status_code == 200
    for _ in range(400):
        if client.get("/api/job").json()["status"] != "running":
            break
        time.sleep(0.05)
    job = client.get("/api/job").json()
    assert job["status"] == "done", job
    assert client.get("/api/preview.jpg", params={"picture": 0, "t": 5}).status_code == 200
    check = client.get("/api/check").json()
    for it in check["items"]:
        client.post("/api/answer", json={"id": it["id"], "answer": "yes"})
    client.post("/api/watch", json={"mode": "skipped"})
    client.post("/api/store", json={"name": "Lismore", "code": "SYN-1"})
    made = client.post("/api/report")
    assert made.status_code == 200, made.text
    r = client.get("/api/report.pptx")
    assert r.status_code == 200 and r.headers["content-type"] == PPTX
    assert client.get("/video", headers={"Range": "bytes=0-9"}).status_code == 206

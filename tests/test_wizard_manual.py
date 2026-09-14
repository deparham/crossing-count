"""The wizard's manual count, and the learning examples both kinds of count leave behind."""

from __future__ import annotations

import json
import time
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from crossing_count.webapp import Setup
from crossing_count.wizard import Wizard, WizardError
from crossing_count.wizard_app import create_wizard_app


def texts(path: Path) -> list[str]:
    return ["\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            for slide in Presentation(str(path)).slides]


@pytest.fixture
def manual(two_tile_video: dict[str, Any], tmp_path: Path) -> Wizard:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("manual")
    w.set_marks(True)
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": False}])
    w.set_direction("in")
    return w


def test_cameras_are_named_without_drawing(manual: Wizard) -> None:
    (cam,) = manual.state["cameras"]
    assert cam["sensor"] == "CN-9-PB1" and cam["site"] == "SYN"
    assert cam["config"] and manual.marked()  # kept for the line's position, but not drawn
    assert manual.state["store"]["code"] == "SYN"
    with pytest.raises(WizardError, match="Name the camera"):
        manual.set_named_cameras([], [{"x0": 0}], [{"picture": 0, "name": " "}])
    with pytest.raises(WizardError, match="nothing to run"):
        manual.start()


def test_footage_from_retailnext_takes_its_store_and_cameras(manual: Wizard) -> None:
    manual.manual_add("CN-9-PB1", 21.0, "in")
    manual.manual_watched("CN-9-PB1", 0.0, 60.0)
    notes = manual.use_download({"subscription": "gazman", "code": "392",
                                 "name": "392 Perri Cutten Armadale",
                                 "cameras": ["Armadale_Entrance", "Back_Door"]})
    store = manual.state["store"]
    assert (store["code"], store["name"]) == ("392", "Perri Cutten Armadale")  # no code twice
    (cam,) = manual.state["cameras"]
    assert (cam["sensor"], cam["site"], cam["config"]) == ("Armadale_Entrance", "392", None)
    assert [c["camera"] for c in manual.state["manual"]["counts"]] == ["Armadale_Entrance"]
    assert "Armadale_Entrance" in manual.state["manual"]["watched"]  # the counts kept, moved
    assert "a camera of store SYN" in notes[-1]
    assert manual.state["decisions"][-1]["action"] == "corrected"
    assert manual.use_download({"code": "392", "cameras": ["Armadale_Entrance"]}) == []


def test_only_the_footages_own_store_drawings_are_offered(two_tile_video: dict[str, Any]) -> None:
    s = Setup(two_tile_video["video"], two_tile_video["dir"])
    assert {c["site"] for c in s.existing()} == {"SYN"}
    s.store = ("392", "392 Perri Cutten Armadale")  # footage known to be from store 392
    assert s.existing() == []
    s.store = ("syn", "")  # its own store's drawings, however the site was written
    assert {c["site"] for c in s.existing()} == {"SYN"}


def test_a_renamed_camera_keeps_its_hand_counts(manual: Wizard) -> None:
    manual.manual_add("CN-9-PB1", 21.0, "in")
    manual.manual_watched("CN-9-PB1", 0.0, 30.0)
    tiles = [c["tile"] for c in manual.state["cameras"]]
    manual.set_named_cameras([], tiles, [{"picture": 0, "name": "Armadale_Entrance"}])
    assert [c["camera"] for c in manual.state["manual"]["counts"]] == ["Armadale_Entrance"]
    assert list(manual.state["manual"]["watched"]) == ["Armadale_Entrance"]


def test_counts_left_under_an_old_name_come_back(manual: Wizard) -> None:
    manual.manual_add("CN-9-PB1", 21.0, "in")
    manual.manual_watched("CN-9-PB1", 0.0, 30.0)
    manual.state["cameras"][0].update(sensor="Armadale_Entrance")  # renamed by an older version
    assert manual.tidy_on_open(pictures=2) == []  # two pictures: whose counts is not known
    notes = manual.tidy_on_open(pictures=1)
    assert notes == [("1 hand count(s) made when this camera was called CN-9-PB1 count again, "
                      "as Armadale_Entrance.")]
    assert manual.manual_summary()["cameras"][0]["in"] == 1
    assert manual.state["decisions"][-1]["action"] == "reattached"


def test_hand_count_and_report(manual: Wizard, tmp_path: Path) -> None:
    with pytest.raises(WizardError, match="Traffic In"):
        manual.manual_add("CN-9-PB1", 10.0, "out")  # this count is for traffic in only
    for t in (21.0, 22.0, 40.0):
        manual.manual_add("CN-9-PB1", t, "in")
    removed = manual.manual_undo("CN-9-PB1")
    assert removed is not None and removed["t"] == 40.0
    manual.manual_watched("CN-9-PB1", 0.0, 30.0)
    part = manual.counts()
    assert part["status"] == "incomplete" and part["unwatched_s"] == 30.0
    assert "Only 50% of CN-9-PB1's footage was watched." in part["incomplete"]
    manual.manual_watched("CN-9-PB1", 29.9, 60.0)
    (cam,) = manual.manual_summary()["cameras"]
    assert cam["in"] == 2 and cam["watched_pct"] > 99 and not cam["unwatched"]
    manual.set_sensor({"in": 2})
    manual.set_store(name="Lismore", code="SYN-1", operator="Pat")
    with pytest.raises(WizardError, match="Finish counting"):
        manual.make_report()
    manual.manual_done()
    c = manual.counts()
    assert c["verified"] == {"in": 2} and c["accuracy"]["in"]["accuracy_pct"] == 100.0
    assert c["status"] == "complete"
    pages = texts(manual.make_report())
    assert "MANUAL COUNT" in pages[0] and "BUSIEST MOMENT" in pages[0]
    assert "Counted by hand" in "\n".join(pages)

    examples = tmp_path / "shared"
    manual.save_examples(examples, wait=True)
    assert manual.state["examples"]["status"] == "saved", manual.state["examples"]
    index = [json.loads(line) for line in (examples / "index.jsonl").read_text().splitlines()]
    labels = [e["label"] for e in index]
    assert labels.count("crossing") == 2 and labels.count("no_crossing") >= 1
    meta = json.loads((examples / index[0]["path"] / "meta.json").read_text())
    assert meta["source"] == "manual_click" and meta["operator"] == "Pat" and meta["frames"]
    assert (examples / index[0]["path"] / meta["frames"][0]["file"]).is_file()
    assert meta["line"] and all(0 <= x <= 640 and 0 <= y <= 480 for x, y in meta["line"])
    (examples / "index.jsonl").write_text(  # another video's entry, which must survive
        (examples / "index.jsonl").read_text() + '{"path": "other/video/cam/crossing_0001"}\n')
    manual.save_examples(examples, wait=True)  # saving again replaces this video's entries
    again = (examples / "index.jsonl").read_text().splitlines()
    assert len(again) == len(index) + 1 and again[-1].startswith('{"path": "syn-1/')
    assert '{"path": "other/video/cam/crossing_0001"}' in again


def test_manual_count_through_the_page_api(two_tile_video: dict[str, Any], tmp_path: Path,
                                           monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path / "home"))
    client = TestClient(create_wizard_app(two_tile_video["dir"], tmp_path,
                                          [two_tile_video["dir"]]))
    client.post("/api/open", json={"path": str(two_tile_video["video"])})
    assert client.get("/api/drawn").json()["marks_guess"] is True  # drawings match the line
    client.post("/api/mode", json={"mode": "manual"})
    client.post("/api/marks", json={"marks": True})
    r = client.post("/api/named-cameras", json={"cameras": [
        {"picture": 0, "name": "CN-9-PB1"}, {"picture": 1, "name": "CN-9-R2"}]})
    assert r.status_code == 200 and len(r.json()["cameras"]) == 2
    client.post("/api/direction", json={"direction": "both"})
    for cam, t, d in (("CN-9-PB1", 21.0, "in"), ("CN-9-R2", 31.0, "in"), ("CN-9-R2", 45.0, "out")):
        assert client.post("/api/hand/add",
                           json={"camera": cam, "t": t, "direction": d}).status_code == 200
    assert client.post("/api/hand/undo", json={"camera": "CN-9-R2"}).json()["removed"]["t"] == 45.0
    for cam in ("CN-9-PB1", "CN-9-R2"):
        client.post("/api/hand/watched", json={"camera": cam, "start": 0, "end": 60})
    hand = client.post("/api/hand/done", json={"done": True}).json()
    assert [c["watched_pct"] for c in hand["hand"]["cameras"]] == [100.0, 100.0]
    assert client.get("/api/hand").json()["geometry"]["CN-9-PB1"]["line"]
    assert client.post("/api/settings", json={"examples_dir": "relative/folder"}).status_code == 400
    shared = str(tmp_path / "shared")
    assert client.post("/api/settings", json={"examples_dir": shared}).json()["examples_dir"] == shared
    client.post("/api/sensor", json={"values": {"in": 2, "out": 0}})
    client.post("/api/store", json={"name": "Lismore", "code": "CN-9", "operator": "Pat"})
    made = client.post("/api/report")
    assert made.status_code == 200, made.text
    ex: dict[str, Any] = {}
    for _ in range(400):
        ex = client.get("/api/examples").json()["examples"]
        if ex["status"] != "saving":
            break
        time.sleep(0.05)
    assert ex["status"] == "saved" and ex["positives"] == 2, ex

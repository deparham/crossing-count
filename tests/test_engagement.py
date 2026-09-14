"""No percentage on too few crossings, and engagements: several finalised windows of one
store put together, the route to a percentage with a range."""

from __future__ import annotations

import hashlib
import json
import stat
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from crossing_count import engagement, runs
from crossing_count import validation as v
from crossing_count.webapp import Setup
from crossing_count.wizard import Wizard
from crossing_count.wizard_app import create_wizard_app


def test_eleven_against_nine_gives_the_counts_not_a_percentage() -> None:
    q = v.quote(11, 9)  # Perri Cutten Armadale 392, 13/09/2026: printed as "81.8%"
    assert (q["rate"], q["error"], q["words"]) == (False, -2, "an undercount of 2")
    assert "one crossing moves it by 9 points" in q["reason"]
    assert v.quote(43, 40)["rate"] and v.quote(43, 43)["words"] == "no difference"
    lv = v.by_level([{"interval": "a", "direction": "in", "truth": 11, "system": 9,
                      "covered_s": 900, "cameras": 1}])
    text = v.level_sentence(lv, "RetailNext")
    assert text is not None and "2 fewer (too few crossings for a percentage)" in text
    assert "%" not in text


def _hand_count(video: dict[str, Any], tmp_path: Path, both: bool = False) -> Wizard:
    w = Wizard(video["video"], video["dir"], tmp_path)
    w.set_mode("manual")
    w.set_marks(True)
    s = Setup(w.video, video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": both}])
    w.state["marks"] = False
    w.set_direction("in")
    w.set_store(name="Armadale", code="CN-9", operator="Alex")
    return w


def _texts(path: Path) -> list[str]:
    return ["\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            for slide in Presentation(str(path)).slides]


def test_the_report_gives_the_difference_where_the_percentage_would_be(
        two_tile_video: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path))
    w = _hand_count(two_tile_video, tmp_path)
    for k in range(11):
        w.manual_add("CN-9-PB1", 2.0 + 5 * k, "in")
    w.manual_watched("CN-9-PB1", 0.0, float(w.state["duration_s"]))
    w.manual_done()
    w.set_sensor({"in": 9})
    c = w.counts()
    assert c["accuracy"]["in"]["accuracy_pct"] == 81.8 and not c["quote"]["in"]["rate"]
    cover = _texts(w.make_report(None))[0]
    assert "81.8" not in cover and "DIFFERENCE" in cover and "−2" in cover
    assert "verified 11, RetailNext counted 9: an undercount of 2" in cover


def test_per_camera_rows_give_no_percentage_on_few_crossings(
        two_tile_video: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path))
    w = _hand_count(two_tile_video, tmp_path, both=True)
    w.manual_add("CN-9-PB1", 5.0, "in")
    w.manual_add("CN-9-R2", 9.0, "in")
    w.set_sensor({"in": 3}, cameras={"CN-9-PB1": {"in": 2}, "CN-9-R2": {"in": 1}})
    b = w._breakdown()
    pb1 = next(c for c in b["cameras"] if c["label"] == "CN-9-PB1")
    assert pb1["accuracy"]["in"]["error"] == 1 and pb1["accuracy"]["in"]["error_pct"] is None
    assert any("only where at least 30" in n for n in b["notes"])


def _run(root: Path, n: int, store: str, truth: int, system: int, role: str,
         supersedes: str | None = None, marked: bool = False) -> str:
    vid = f"CC-VAL-2026-ABCD-{n:06d}"
    folder = runs.runs_dir(root) / vid
    folder.mkdir(parents=True)
    result = {"status": "complete", "marked": marked, "system": "RetailNext",
              "store": {"code": store}, "cameras": ["C1"], "duration_s": 900, "fingerprint": vid,
              "sampling": {"role": role},
              "rows": [{"interval": "whole footage", "direction": "out", "truth": truth,
                        "system": system, "covered_s": 900, "cameras": 1}]}
    text = json.dumps(result)
    (folder / "result.json").write_text(text)
    (folder / "manifest.json").write_text(json.dumps({
        "schema": runs.SCHEMA, "validation_id": vid, "supersedes": supersedes,
        "store": {"code": store, "name": "Armadale"}, "footage": {"clock_start": None},
        "audit_log_head": None,
        "files": {"result.json": hashlib.sha256(text.encode()).hexdigest()}}))
    return vid


def test_windows_put_together_give_a_percentage_and_its_range(tmp_path: Path) -> None:
    few = [_run(tmp_path, k, "392", 11 + k, 9 + k, "peak") for k in range(1, 3)]  # 25 verified
    e = engagement.summarise(tmp_path, few)
    assert not e["directions"]["out"]["rate"] and e["directions"]["out"]["bias_pct"] is None
    assert "RetailNext counted 21: an undercount of 4" in e["headline"]
    three = [*few, _run(tmp_path, 3, "392", 11, 9, "peak")]  # 36 verified: a percentage
    out = engagement.summarise(tmp_path, three)["directions"]["out"]
    assert out["rate"] and out["bias_pct"] == -16.7 and out["range"] is None
    assert out["range_note"] == "no range: fewer than 5 windows"
    peaks = [_run(tmp_path, 10 + k, "CN-1", 40, 36, "peak") for k in range(4)]  # 160 an hour
    controls = [_run(tmp_path, 20 + k, "CN-1", 20, 20, "control") for k in range(2)]  # 80
    e = engagement.summarise(tmp_path, [*peaks, *controls])
    total = e["directions"]["out"]
    assert (total["truth"], total["system"], total["bias_pct"]) == (200, 184, -8.0)
    assert total["range"] is not None and total["range"][0] <= -8.0 <= total["range"][1]
    assert e["summary"]["scope"] == {"peak": 4, "control": 2}
    assert e["levels"] is not None and "busy traffic" in e["levels"] and "10.0% too few" in e["levels"]
    with pytest.raises(engagement.EngagementError, match="one store's windows"):
        engagement.summarise(tmp_path, [peaks[0], few[0]])
    marked = _run(tmp_path, 30, "CN-1", 40, 40, "peak", marked=True)
    newer = _run(tmp_path, 31, "CN-1", 40, 38, "peak", supersedes=peaks[0])
    e = engagement.summarise(tmp_path, [*peaks, marked, newer])
    assert {x["id"]: x["why"][0] for x in e["left"]} == {
        peaks[0]: f"replaced by {newer}, also chosen",
        marked: "counted on footage showing the system's own marks"}
    kept = engagement.write(tmp_path, [*peaks, *controls], by="Alex")
    p = Path(kept["folder"]) / "engagement.json"
    rec = json.loads(p.read_text())
    assert rec["engagement_id"].startswith("CC-ENG-") and len(rec["validations"]) == 6
    assert not p.stat().st_mode & stat.S_IWUSR


def test_the_runs_page_puts_windows_together(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CROSSING_COUNT_HOME", str(tmp_path))
    ids = [_run(tmp_path, k, "CN-1", 40, 36, "peak") for k in range(1, 3)]
    client = TestClient(create_wizard_app(tmp_path / "sites", tmp_path, folders=[tmp_path]))
    got = client.post("/api/engagement", json={"ids": ids})
    assert got.status_code == 200 and got.json()["directions"]["out"]["bias_pct"] == -10.0
    assert client.post("/api/engagement", json={"ids": []}).status_code == 400

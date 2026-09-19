"""Counting the total only: how many crossed, without saying when any of them did.

The point of this mode is that it needs none of the rest: no detector, no tracker, no
recorded detections, no gold clip, no matching. The point of these tests is that it also
never reaches any of them - a total saved here must never turn into a crossing the tool
can be scored against.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from crossing_count import bench, gold
from crossing_count.webapp import Setup
from crossing_count.wizard import Wizard, WizardError
from crossing_count.wizard_app import create_wizard_app


def _report_text(path: Path) -> list[str]:
    from pptx import Presentation
    return ["\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            for slide in Presentation(str(path)).slides]


@pytest.fixture
def totals(two_tile_video: dict[str, Any], tmp_path: Path) -> Wizard:
    """Footage open, cameras named, traffic chosen: nothing else, and nothing has run."""
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("total")
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": True}])
    w.set_direction("in")
    return w


def test_pressing_up_and_down(totals: Wizard) -> None:
    assert totals.total_only() and not totals.manual()
    assert totals.total_by_direction() == {"in": 0}
    assert totals.total_step("CN-9-PB1", "in") == 1
    assert totals.total_step("CN-9-PB1", "in") == 2
    assert totals.total_step("CN-9-PB1", "in", -1) == 1
    assert totals.total_step("CN-9-PB1", "in", -1) == 0
    assert totals.total_step("CN-9-PB1", "in", -1) == 0  # never below nobody
    with pytest.raises(WizardError, match="one person up or down"):
        totals.total_step("CN-9-PB1", "in", 5)
    with pytest.raises(WizardError, match="This count is for"):
        totals.total_step("CN-9-PB1", "out")
    with pytest.raises(WizardError, match="not one of this count's cameras"):
        totals.total_step("CN-9-NOPE", "in")
    # what was pressed is kept, so a total that looks wrong later can be read back
    assert [a["step"] for a in totals.state["totals"]["actions"]] == [1, 1, -1, -1]
    assert all("t" not in a for a in totals.state["totals"]["actions"])  # never a moment


def test_each_camera_and_direction_keeps_its_own_number(
        two_tile_video: dict[str, Any], tmp_path: Path) -> None:
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("total")
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": True}])
    w.set_direction("both")
    w.total_set("CN-9-PB1", "in", 100)
    w.total_set("CN-9-PB1", "out", 40)
    w.total_set("CN-9-R2", "in", 47)
    assert w.total_by_direction() == {"in": 147, "out": 40}
    assert w.total_missing() == ["Enter Traffic Out for CN-9-R2."]
    w.total_set("CN-9-R2", "out", 0)
    assert w.total_missing() == []
    with pytest.raises(WizardError, match="cannot be negative"):
        w.total_set("CN-9-R2", "out", -1)


def test_a_total_is_finished_only_when_it_covers_what_it_is_compared_with(totals: Wizard) -> None:
    with pytest.raises(WizardError, match="Enter Traffic In for CN-9-PB1"):
        totals.total_done()
    totals.total_set("CN-9-PB1", "in", 100)
    totals.total_set("CN-9-R2", "in", 47)
    with pytest.raises(WizardError, match="whole clip"):
        totals.total_done()
    totals.total_notes("Busy period, several overlapping people.")
    totals.total_done(whole_clip=True)
    got = totals.total_summary()
    assert got["done"] and got["by_direction"] == {"in": 147}
    assert got["notes"] == "Busy period, several overlapping people."
    assert totals.counts()["status"] == "complete"
    assert [c["ok"] for c in totals.counts()["checks"][:3]] == [True, True, True]
    totals.total_step("CN-9-PB1", "in")  # changed afterwards: no longer finished
    assert not totals.total_summary()["done"]
    assert totals.counts()["status"] == "incomplete"


def test_a_saved_total_never_becomes_a_crossing(totals: Wizard) -> None:
    """The separation test. Counting 147 people says nothing about when any of them crossed,
    so nothing here may be matched, scored, or kept as ground truth."""
    totals.total_set("CN-9-PB1", "in", 100)
    totals.total_set("CN-9-R2", "in", 47)
    totals.total_done(whole_clip=True)

    assert totals.verified_rows() == []  # no crossing has a moment
    assert totals.unsure_rows() == []
    c = totals.counts()
    assert c["verified"] == {"in": 147} and c["detected"] == {"in": 0}
    assert c["found"] == 0 and c["not_found"] == 0  # no match was made, right or wrong
    assert c["total_only"] is True
    # nothing to run, and nothing the benchmark or the gold set will take
    with pytest.raises(WizardError, match="nothing to run"):
        totals.start()
    assert bench.find_truth(totals.run_dir) is None
    (why,) = gold.problems(totals.state)
    assert "total only" in why and "count by hand that marks every crossing" in why
    assert totals.result()["specification"] is None  # no version of what a crossing is applied


def test_the_total_is_compared_with_the_system_as_one_number_against_another(
        totals: Wizard) -> None:
    totals.total_set("CN-9-PB1", "in", 100)
    totals.total_set("CN-9-R2", "in", 47)
    totals.total_done(whole_clip=True)
    totals.set_sensor({"in": 139})

    comp = totals.comparison()
    assert comp["total_only"] and comp["verified"] == {"in": 147}
    assert comp["intervals"] == []  # a total belongs to no one 15-minute interval
    (row,) = comp["rows"]
    assert row["truth"] == 147 and row["system"] == 139
    assert row["interval"] == "whole footage" and row["reference"] == "total only"
    got = comp["metrics"]["in"]
    assert got["truth"] == 147 and got["system"] == 139
    assert got["error"] == -8 and got["undercount"] == 8 and got["overcount"] == 0
    assert got["wape_pct"] == pytest.approx(100 * 8 / 147, abs=0.05)  # 5.44%
    assert got["intervals"] == 1  # one number counted, one number to set it against
    assert [c["counts"]["in"] for c in totals.total_summary()["cameras"]] == [100, 47]


def test_the_whole_way_through_without_the_automatic_side(
        two_tile_video: dict[str, Any], tmp_path: Path) -> None:
    """Video, total, save, report - over the web pages, with nothing else available."""
    client = TestClient(create_wizard_app(two_tile_video["dir"], tmp_path,
                                          folders=[two_tile_video["video"].parent]))
    assert client.post("/api/open", json={"path": str(two_tile_video["video"])}).status_code == 200
    assert client.post("/api/mode", json={"mode": "total"}).status_code == 200
    assert client.post("/api/named-cameras", json={"cameras": [
        {"picture": 0, "name": "CN-9-PB1", "include": True},
        {"picture": 1, "name": "CN-9-R2", "include": False}]}).status_code == 200
    assert client.post("/api/direction", json={"direction": "in"}).status_code == 200
    for _ in range(3):
        got = client.post("/api/total/step",
                          json={"camera": "CN-9-PB1", "direction": "in"}).json()
    assert got["total"]["by_direction"] == {"in": 3}
    client.post("/api/total/step", json={"camera": "CN-9-PB1", "direction": "in", "step": -1})
    client.post("/api/total/notes", json={"notes": "Quiet, one pram."})
    done = client.post("/api/total/done", json={"done": True, "whole_clip": True})
    assert done.status_code == 200 and done.json()["total"]["done"]
    assert client.get("/api/total").json()["total"]["by_direction"] == {"in": 2}
    # the count survives the page being reopened
    again = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    assert again.total_by_direction() == {"in": 2} and again.total_summary()["done"]
    assert again.total_summary()["notes"] == "Quiet, one pram."
    # nothing was detected, tracked or scored to get here
    assert again.state["job"]["status"] == "idle"
    assert not list((again.run_dir).rglob("detections.pkl"))
    assert not list((again.run_dir).rglob("candidates.json"))


def test_a_count_by_hand_is_unaffected(two_tile_video: dict[str, Any], tmp_path: Path) -> None:
    """Changing modes must not let one kind of count be read as the other."""
    w = Wizard(two_tile_video["video"], two_tile_video["dir"], tmp_path)
    w.set_mode("total")
    s = Setup(w.video, two_tile_video["dir"])
    w.set_named_cameras(s.existing(), [t.as_dict() for t in s.tiles],
                        [{"picture": 0, "name": "CN-9-PB1", "include": True},
                         {"picture": 1, "name": "CN-9-R2", "include": False}])
    w.set_direction("in")
    w.total_set("CN-9-PB1", "in", 147)
    w.set_mode("manual")  # the same footage, counted properly this time
    assert w.verified_rows() == [] and w.counts()["verified"] == {"in": 0}
    w.manual_add("CN-9-PB1", 1.0, "in")
    assert w.counts()["verified"] == {"in": 1}  # the total is not added in
    assert gold.problems(w.state)[0] != gold.problems({**w.state, "mode": "total"})[0]
    w.set_mode("total")
    assert w.total_by_direction() == {"in": 147}  # and it was kept, not thrown away


def test_choosing_a_mode_that_does_not_exist(totals: Wizard) -> None:
    with pytest.raises(WizardError, match="automatic counting, counting by hand, or the total"):
        totals.set_mode("guess")


def test_the_counting_page_has_what_it_needs(two_tile_video: dict[str, Any],
                                             tmp_path: Path) -> None:
    """The page asks for the pictures' geometry and the total: both must answer in this mode."""
    client = TestClient(create_wizard_app(two_tile_video["dir"], tmp_path,
                                          folders=[two_tile_video["video"].parent]))
    client.post("/api/open", json={"path": str(two_tile_video["video"])})
    client.post("/api/mode", json={"mode": "total"})
    client.post("/api/named-cameras", json={"cameras": [
        {"picture": 0, "name": "CN-9-PB1", "include": True},
        {"picture": 1, "name": "CN-9-R2", "include": True}]})
    client.post("/api/direction", json={"direction": "in"})
    hand = client.get("/api/hand")
    assert hand.status_code == 200 and set(hand.json()["geometry"]) == {"CN-9-PB1", "CN-9-R2"}
    page = client.get("/api/total").json()
    assert page["mode"] == "total" and page["dirs"] == ["in"]
    assert [c["sensor"] for c in page["total"]["cameras"]] == ["CN-9-PB1", "CN-9-R2"]
    assert page["totals"]["done"] is False  # what the page reads to know where it is


def test_the_report_says_what_a_total_can_and_cannot_show(totals: Wizard) -> None:
    totals.total_set("CN-9-PB1", "in", 100)
    totals.total_set("CN-9-R2", "in", 47)
    totals.total_notes("Busy period, several overlapping people.")
    totals.total_done(whole_clip=True)
    totals.set_sensor({"in": 139})
    totals.set_store(name="Lismore", code="CN-9", operator="Sam")

    data = totals.report_data(totals.counts(), [], Path("frame.jpg"), "", [])
    method = " ".join(data["method"])
    assert "pressed a key for each person, keeping the number only" in method
    assert "The moment each person crossed was not recorded" in method
    assert "Busy period, several overlapping people." in method
    assert "one total against another" in method
    assert "an over-count and an under-count in the same period cancel out" in method
    assert "no recall, precision or missed-crossing rate can be worked out" in method
    # no version of what a crossing is, because no crossing was marked
    assert "Ground Truth Specification" not in method
    assert "Who was counted: children counted, staff counted." in method
    assert "100 traffic in" in method and "47 traffic in" in method


def test_video_to_total_to_saved_count_to_report(totals: Wizard) -> None:
    """The whole way through, with no detector, no tracker and no recorded detections."""
    for _ in range(100):
        totals.total_step("CN-9-PB1", "in")
    totals.total_set("CN-9-R2", "in", 47)
    totals.total_notes("Busy period, several overlapping people.")
    totals.total_done(whole_clip=True)
    totals.set_sensor({"in": 139})
    totals.set_store(name="Lismore", code="CN-9", operator="Sam")
    assert totals.counts()["status"] == "complete"

    pages = "\n".join(t for t in _report_text(totals.make_report()))
    assert "MANUAL COUNT" in pages and "VERIFIED COUNT" not in pages
    assert "Total only: no moment recorded for any crossing" in pages
    # the page that lists the crossings says why there are none, rather than "none verified"
    assert "Counted as a total only: 147 people coming in" in pages
    assert "No crossings were verified for this period" not in pages
    assert "keeping the number only" in pages
    assert "no recall, precision or missed-crossing rate can be worked out" in pages
    assert "147" in pages and "139" in pages


def test_nothing_of_the_tools_is_claimed_on_a_total(totals: Wizard) -> None:
    """Everywhere that asks "was this counted by hand?" to decide what the tool did must
    treat a total as a count by a person, not as a check of an automatic count."""
    totals.total_set("CN-9-PB1", "in", 100)
    totals.total_set("CN-9-R2", "in", 47)
    totals.total_done(whole_clip=True)
    totals.set_sensor({"in": 139})
    totals.set_store(name="Lismore", code="CN-9", operator="Sam")

    data = totals.report_data(totals.counts(), [], Path("frame.jpg"), "", [])
    assert "AUTOMATED DETECTION OVERLAY" not in data["frames_title"]
    assert totals.comparison()["overlap"] == 0  # no proposals, so no camera overlap to work out
    assert totals.unsure_rows() == []
    frame, caption = totals.render_busy_frame()  # the tool never ran: no overlay to draw
    assert frame.is_file() and "CN-9-PB1" in caption
    assert totals.tool_vs_hand() is None

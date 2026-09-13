"""The validation report as a PowerPoint file, laid out like iTOi's camera validation page.

Page 1 follows the original: logo and date, "<store> <code> Traffic System",
the CAMERA VALIDATION tag, the capture details, three cards (verified count,
system count, accuracy) and a detection-overlay frame. The pages after it hold
the details: how the count was made, every verified crossing with its time,
and a close-up of each one. A4 portrait, like the PDF.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

A4_W, A4_H = Emu(7560000), Emu(10692000)  # 210 x 297 mm
NAVY = RGBColor(0x10, 0x1D, 0x3B)
TEAL = RGBColor(0x1F, 0x9B, 0xB3)
TEAL_TEXT = RGBColor(0x23, 0x9A, 0xB4)
PILL_BG = RGBColor(0xE8, 0xF2, 0xF4)
PILL_TEXT = RGBColor(0x1C, 0x6B, 0x78)
CARD_BG = RGBColor(0xEE, 0xF4, 0xF6)
GREY = RGBColor(0x4A, 0x55, 0x68)
ROW_BG = RGBColor(0xF5, 0xF8, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN_TEXT = RGBColor(0xB0, 0x3A, 0x2E)
SERIF = "Georgia"
SANS = "Arial"  # installed on every Mac and PC; Calibri is not on Macs without Office
FOOTER = "iTOi Solutions  –  Confidential"
TABLE_TOP_MAX = 10.55  # inches: tables and pictures stay above the footer


def _box(slide: Any, left: float, top: float, width: float, height: float) -> Any:
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb


def _write(shape: Any, text: str, size: float, *, bold: bool = False, color: RGBColor = NAVY,
           font: str = SANS, align: Any = PP_ALIGN.LEFT, spacing: float | None = None,
           anchor: Any = None) -> None:
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    if spacing:  # letter spacing, in hundredths of a point
        run._r.get_or_add_rPr().set("spc", str(round(spacing * 100)))
    if anchor is not None:
        tf.vertical_anchor = anchor


def _shadow(shape: Any) -> None:
    effects = etree.SubElement(shape._element.spPr, qn("a:effectLst"))
    shadow = etree.SubElement(effects, qn("a:outerShdw"), blurRad="88900", dist="25400",
                              dir="5400000", algn="t", rotWithShape="0")
    colour = etree.SubElement(shadow, qn("a:srgbClr"), val="000000")
    etree.SubElement(colour, qn("a:alpha"), val="16000")


def _panel(slide: Any, left: float, top: float, width: float, height: float, fill: RGBColor,
           radius: float = 0.07, shadow: bool = False) -> Any:
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                               Inches(width), Inches(height))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if shadow:
        _shadow(s)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    return s


def _picture(slide: Any, path: str, left: float, top: float, max_w: float,
             max_h: float) -> tuple[float, float]:
    """Add a picture fitted inside max_w x max_h, centred horizontally; return its size."""
    with PILImage.open(path) as im:
        w_px, h_px = im.size
    scale = min(max_w / w_px, max_h / h_px)
    w, h = w_px * scale, h_px * scale
    slide.shapes.add_picture(path, Inches(left + (max_w - w) / 2), Inches(top), Inches(w),
                             Inches(h))
    return w, h


def _pct(v: float | None) -> str:
    if v is None:
        return "–"
    return f"{v:.0f}%" if float(v).is_integer() else f"{v:.1f}%"


def _value_size(value: str, base: float) -> float:
    """Card numbers at full size, a range a little smaller, a word (INCOMPLETE) smaller still."""
    if value[:1].isalpha():
        return round(base * 0.4, 1)
    return base if len(value) <= 5 else round(base * 5 / len(value), 1)


def _header(slide: Any, data: dict[str, Any], logo: Path | None) -> None:
    if logo is not None and logo.is_file():
        _picture(slide, str(logo), 0.52, 0.26, 2.0, 0.74)
    else:
        _write(_box(slide, 0.6, 0.38, 3.0, 0.5), "i to i solutions — image to intelligence", 13,
               bold=True, color=NAVY)
    _write(_box(slide, 4.0, 0.40, 3.64, 0.3), "OPTIMIZATION REPORT", 12, bold=True,
           color=TEAL_TEXT, align=PP_ALIGN.RIGHT, spacing=3)
    _write(_box(slide, 4.0, 0.68, 3.64, 0.3), data["report_date"], 14, color=GREY,
           align=PP_ALIGN.RIGHT)


def _footer(slide: Any) -> Any:
    _write(_box(slide, 0.63, 11.05, 4.0, 0.25), FOOTER, 10, color=GREY)
    return _box(slide, 4.6, 11.05, 3.04, 0.25)  # "Page k of N", written once all pages exist


def _cover(slide: Any, data: dict[str, Any]) -> None:
    title = f"{data['store_name']} {data['store_code']} Traffic System"
    size = 28.0 if len(title) <= 28 else max(17.0, 28.0 * 28 / len(title))  # stay on one line
    _write(_box(slide, 0.62, 1.27, 7.1, 0.6), title, size, bold=True, font=SERIF)
    pill = _panel(slide, 0.53, 1.93, 2.37, 0.36, PILL_BG, radius=0.2)
    pill.text_frame.margin_left = pill.text_frame.margin_right = Inches(0.04)
    pill.text_frame.word_wrap = False
    _write(pill, "CAMERA VALIDATION", 10, bold=True, color=PILL_TEXT, align=PP_ALIGN.CENTER,
           spacing=1.6, anchor=MSO_ANCHOR.MIDDLE)
    _write(_box(slide, 0.62, 2.52, 7.0, 0.5), data["store_code"], 24, bold=True)
    rows = data["directions"]
    captured = f"Captured {data['captured_date']} Cameras {data['location']} {data['time_range']}"
    if len(rows) == 1:
        captured += f"  ·  {rows[0]['label']}"
    _write(_box(slide, 0.62, 3.08, 7.0, 0.35), captured, 14, color=GREY)

    one = len(rows) == 1
    card_h = 1.49 if one else 1.08
    top = 3.70
    kind = str(data.get("count_label", "VERIFIED COUNT")).split()[0]  # VERIFIED or MANUAL
    complete = bool(data.get("complete", True))
    for r in rows:
        tag = "" if one else f" {r['key'].upper()}"
        unsure = int(r.get("unsure") or 0)
        span = r.get("accuracy_range")
        if not complete:  # footage nobody watched may hold people missing from the count
            accuracy = "INCOMPLETE"
        elif unsure and span:
            accuracy = (_pct(span[0]) if span[0] == span[1]
                        else f"{_pct(span[0])[:-1]}–{_pct(span[1])}")
        else:
            accuracy = _pct(r["accuracy"])
        verified = f"{r['verified']}–{r['verified'] + unsure}" if unsure else str(r["verified"])
        cards = ((f"{kind}{tag or ' COUNT'}", verified),
                 (f"SYSTEM{tag or ' COUNT'}", str(r["system"])),
                 (f"ACCURACY{tag}", accuracy))
        for j, (label, value) in enumerate(cards):
            left, hi = 0.53 + j * 2.47, j == 2
            _panel(slide, left, top, 2.24, card_h, TEAL if hi else CARD_BG, shadow=True)
            _write(_box(slide, left, top + (0.27 if one else 0.15), 2.24, 0.3), label, 11,
                   bold=True, color=WHITE if hi else GREY, align=PP_ALIGN.CENTER, spacing=2)
            _write(_box(slide, left, top + (0.58 if one else 0.40), 2.24, 0.8 if one else 0.6),
                   value, _value_size(value, 50 if one else 36), bold=True, font=SERIF,
                   color=WHITE if hi else NAVY, align=PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE if value[:1].isalpha() else MSO_ANCHOR.TOP)
        top += card_h + 0.22
    if not complete:
        note = _box(slide, 0.53, top - 0.08, 7.2, 0.5)
        note.text_frame.word_wrap = True
        _write(note, "Validation incomplete: " + " ".join(data.get("incomplete", []))
               + " No accuracy is given.", 10, bold=True, color=WARN_TEXT)
        top += 0.5

    _write(_box(slide, 0.5, top, 7.27, 0.3),
           str(data.get("frames_title", "VALIDATION FRAMES — AUTOMATED DETECTION OVERLAY")),
           12, bold=True, align=PP_ALIGN.CENTER, spacing=2.5)
    img_top = top + 0.42
    frame = data.get("frame")
    if frame and Path(frame).is_file():
        _, h = _picture(slide, frame, 0.91, img_top, 6.45, TABLE_TOP_MAX - 0.3 - img_top)
        if data.get("frame_caption"):
            _write(_box(slide, 0.5, img_top + h + 0.06, 7.27, 0.25), data["frame_caption"], 9,
                   color=GREY, align=PP_ALIGN.CENTER)


def _cell(cell: Any, text: str, *, header: bool = False, shade: bool = False) -> None:
    cell.text = text
    run = cell.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(10 if header else 9.5)
    run.font.bold = header
    run.font.name = SANS
    run.font.color.rgb = WHITE if header else NAVY
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY if header else (ROW_BG if shade else WHITE)
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _table(slide: Any, rows: list[dict[str, Any]], top: float) -> None:
    cols = (("#", 0.45), ("Time", 1.0), ("Camera", 1.75), ("Direction", 1.1),
            ("How it was found", 2.9))
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(0.53), Inches(top),
                                   Inches(7.2), Inches(0.26 * (len(rows) + 1)))
    tbl = shape.table
    for j, (name, width) in enumerate(cols):
        tbl.columns[j].width = Inches(width)
        _cell(tbl.cell(0, j), name, header=True)
    tbl.rows[0].height = Inches(0.3)
    for i, r in enumerate(rows, 1):
        for j, value in enumerate((str(r["n"]), r["time"], r["camera"], r["direction"],
                                   r["found"])):
            _cell(tbl.cell(i, j), value, shade=i % 2 == 0)
        tbl.rows[i].height = Inches(0.26)


def _details(new_page: Any, data: dict[str, Any]) -> None:
    slide = new_page()
    _write(_box(slide, 0.62, 1.27, 7.1, 0.5), "Crossing details", 24, bold=True, font=SERIF)
    _write(_box(slide, 0.62, 1.85, 7.1, 0.3),
           f"{data['store_name']} {data['store_code']}  ·  Captured {data['captured_date']} "
           f"{data['time_range']}", 12, color=GREY)
    method = data["method"]
    height = 0.3 + 0.42 * len(method)
    panel = _panel(slide, 0.53, 2.3, 7.2, height, CARD_BG)
    tf = panel.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for k, line in enumerate(method):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT  # shapes centre text by default
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(10)
        run.font.name = SANS
        run.font.color.rgb = NAVY
    rows = data["crossings"]
    top = 2.3 + height + 0.3
    if not rows:
        _write(_box(slide, 0.62, top, 7.0, 0.3), "No crossings were verified for this period.",
               12, color=GREY)
        return
    per_row = 0.26
    while rows:
        fit = max(1, int((TABLE_TOP_MAX - top - 0.3) / per_row))
        _table(slide, rows[:fit], top)
        rows = rows[fit:]
        if rows:
            slide = new_page()
            _write(_box(slide, 0.62, 1.27, 7.1, 0.5), "Crossing details (continued)", 24,
                   bold=True, font=SERIF)
            top = 1.95


def _diff(acc: dict[str, Any] | None) -> str:
    """The system minus the verified count, and as a share of it: "+2 (+10%)"."""
    if not acc or acc.get("error") is None:
        return "–"
    err, pct = int(acc["error"]), acc.get("error_pct")
    return f"{err:+d}" if pct is None else f"{err:+d} ({pct:+.0f}%)"


def _grid(slide: Any, header: list[str], rows: list[list[str]], top: float,
          widths: list[float]) -> float:
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.53), Inches(top),
                                   Inches(sum(widths)), Inches(0.3 + 0.26 * len(rows)))
    tbl = shape.table
    for j, (name, width) in enumerate(zip(header, widths, strict=True)):
        tbl.columns[j].width = Inches(width)
        _cell(tbl.cell(0, j), name, header=True)
    tbl.rows[0].height = Inches(0.3)
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            _cell(tbl.cell(i, j), value or "–", shade=i % 2 == 0)
        tbl.rows[i].height = Inches(0.26)
    return top + 0.3 + 0.26 * len(rows)


def _chart(slide: Any, top: float, title: str, labels: list[str], verified: list[int],
           system: list[int | None]) -> float:
    data = CategoryChartData()
    data.categories = labels
    data.add_series("Verified", verified)
    data.add_series("System", system)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.53), Inches(top),
                                   Inches(7.2), Inches(2.35), data).chart
    chart.font.size = Pt(9)
    chart.font.name = SANS
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    for series, colour in zip(chart.series, (TEAL, NAVY), strict=False):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
    return top + 2.5


def _breakdown(new_page: Any, data: dict[str, Any]) -> None:
    """Verified against the system per 15-minute interval and per camera, where there is
    more than one; charts where the system's interval numbers are known."""
    b = data.get("breakdown") or {}
    ivs, cams, dirs = b.get("intervals") or [], b.get("cameras") or [], b.get("dirs") or []
    if not ivs and not cams:
        return
    title = "Traffic by 15 minutes" if ivs else "Traffic by camera"
    slide = new_page()
    _write(_box(slide, 0.62, 1.27, 7.1, 0.5), title, 24, bold=True, font=SERIF)
    top = 1.95

    def room(height: float) -> None:
        nonlocal slide, top
        if top + height > TABLE_TOP_MAX:
            slide = new_page()
            _write(_box(slide, 0.62, 1.27, 7.1, 0.5), f"{title} (continued)", 24, bold=True,
                   font=SERIF)
            top = 1.95

    cols = [f"{d['label'].split()[-1]} {k}" for d in dirs for k in ("verified", "system", "difference")]
    first = 1.55
    widths = [first, *[(7.2 - first) / len(cols)] * len(cols)]

    def rows_of(items: list[dict[str, Any]]) -> list[list[str]]:
        out = []
        for it in items:
            row = [it["label"] + (" (part)" if it.get("partial") else "")]
            for d in dirs:
                system = (it.get("sensor") or {}).get(d["key"])
                row += [str(it["verified"][d["key"]]), "–" if system is None else str(system),
                        _diff((it.get("accuracy") or {}).get(d["key"]))]
            out.append(row)
        return out

    if ivs:
        if any(it.get("sensor") for it in ivs):
            for d in dirs:
                room(2.5)
                top = _chart(slide, top, d["label"], [it["label"].split(" - ")[0] for it in ivs],
                             [it["verified"][d["key"]] for it in ivs],
                             [(it.get("sensor") or {}).get(d["key"]) for it in ivs])
        room(0.6 + 0.26 * len(ivs))
        top = _grid(slide, ["Time", *cols], rows_of(ivs), top, widths) + 0.3
    if cams:
        room(0.6 + 0.26 * len(cams))
        top = _grid(slide, ["Camera", *cols], rows_of(cams), top, widths) + 0.3
    for note in b.get("notes") or []:
        room(0.5)
        box = _box(slide, 0.53, top, 7.2, 0.45)
        box.text_frame.word_wrap = True
        _write(box, note, 10, color=GREY)
        top += 0.5


def _snapshots(new_page: Any, data: dict[str, Any]) -> None:
    thumbs = data["thumbs"]
    per_page, cols, col_w, size = 12, 3, 2.33, 1.85
    for k in range(0, len(thumbs), per_page):
        slide = new_page()
        _write(_box(slide, 0.62, 1.27, 7.1, 0.5), "Crossing snapshots", 24, bold=True,
               font=SERIF)
        for idx, th in enumerate(thumbs[k:k + per_page]):
            row, col = divmod(idx, cols)
            left = 0.64 + col * col_w + 0.05
            top = 1.95 + row * 2.2
            _, h = _picture(slide, th["path"], left, top, col_w - 0.1, size)  # keeps its shape
            _write(_box(slide, left, top + h + 0.04, col_w - 0.1, 0.25), th["label"], 9,
                   color=GREY, align=PP_ALIGN.CENTER)


def build_report(data: dict[str, Any], out: Path, logo: Path | None = None) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = A4_W, A4_H
    blank = prs.slide_layouts[6]
    page_boxes: list[Any] = []

    def new_page() -> Any:
        slide = prs.slides.add_slide(blank)
        _header(slide, data, logo)
        page_boxes.append(_footer(slide))
        return slide

    _cover(new_page(), data)
    _breakdown(new_page, data)
    _details(new_page, data)
    _snapshots(new_page, data)
    for k, box in enumerate(page_boxes, 1):
        _write(box, f"Page {k} of {len(page_boxes)}", 10, color=GREY, align=PP_ALIGN.RIGHT)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
